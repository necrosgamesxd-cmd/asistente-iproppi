# -*- coding: utf-8 -*-
"""Índice de clientes en Excel (openpyxl) — mismo formato que la skill."""
import io

COLUMNAS = [
    ("Nombre", "nombre"),
    ("Estado", "estado"),
    ("Teléfono", "Teléfono"),
    ("Correo", "Correo"),
    ("Renta líquida", "Renta líquida"),
    ("Cap. ahorro", "Capacidad de ahorro"),
    ("Capacidad UF", "Capacidad calculada (UF)"),
    ("Comunas objetivo", "Comunas objetivo"),
    ("Tipo", "Tipo de propiedad"),
    ("Proyecto favorito", "Proyecto favorito"),
    ("Próxima reunión", "Próxima reunión"),
    ("Última reunión", "última"),
    ("Qué busca (resumen)", "busca"),
    ("Carpeta", "carpeta"),
]
ANCHOS = [28, 16, 16, 28, 16, 16, 20, 32, 26, 26, 20, 14, 55, 28]


def construir(filas):
    """-> bytes del archivo clientes.xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append([c for c, _ in COLUMNAS])
    for fila in filas:
        ws.append([fila.get(origen, "") for _, origen in COLUMNAS])

    relleno = PatternFill("solid", fgColor="1F4E78")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")
    for i, ancho in enumerate(ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
